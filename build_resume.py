"""Generate a 1-page Nucleate Chicago resume.docx for Andrew Freiburger."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0B, 0x3D, 0x91)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(9.5)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.space_before = Pt(0)


def set_spacing(p, before=0, after=0, line=1.05):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_run(p, text, *, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return r


def section_header(text):
    p = doc.add_paragraph()
    set_spacing(p, before=3, after=1)
    add_run(p, text.upper(), bold=True, size=10.5, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0B3D91")
    pBdr.append(bottom)
    pPr.append(pBdr)


def entry_line(left_bold, left_rest, right):
    p = doc.add_paragraph()
    set_spacing(p, before=1, after=0)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.4), WD_ALIGN_PARAGRAPH.RIGHT)
    add_run(p, left_bold, bold=True)
    if left_rest:
        add_run(p, left_rest)
    add_run(p, "\t")
    add_run(p, right, italic=True, color=GREY)
    return p


def bullet_runs(text_runs, indent=0.18):
    p = doc.add_paragraph(style="Normal")
    set_spacing(p, before=0, after=0)
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    add_run(p, "• ")
    for t, opts in text_runs:
        add_run(p, t, bold=opts.get("bold", False), italic=opts.get("italic", False))
    return p


# ===================== HEADER =====================
name_p = doc.add_paragraph()
set_spacing(name_p, before=0, after=0)
name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(name_p, "ANDREW P. FREIBURGER", bold=True, size=18, color=NAVY)

tag = doc.add_paragraph()
set_spacing(tag, before=1, after=1)
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(tag, "Computational biologist building metabolic-modeling platforms for microbiome therapeutics", size=10, italic=True, color=NAVY)

sub = doc.add_paragraph()
set_spacing(sub, before=0, after=1)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(sub, "PhD Candidate, Chemical & Biological Engineering  ·  Northwestern University", size=9, color=GREY)

links = doc.add_paragraph()
set_spacing(links, before=0, after=3)
links.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(links, "andrewfreiburger2027@u.northwestern.edu  ·  andrewfreiburger.com  ·  github.com/freiburgermsu", size=9)

# ===================== EDUCATION =====================
section_header("Education")
entry_line("Northwestern University", " — Ph.D., Chemical and Biological Engineering  ·  Tyo Lab", "Sep 2023 – Present")
entry_line("University of Victoria", " — M.A.Sc., Civil Engineering  ·  GPA 8.0/9.0 (4.0 eqv.)  ·  3× UVic Graduate Award", "Jan 2020 – Apr 2022")
entry_line("Grand Valley State University", " — B.S. (honors), Chemistry; Biology minor; Green Chemistry cert.  ·  GPA 3.6/4.0", "Aug 2016 – Apr 2019")
p = doc.add_paragraph()
set_spacing(p, before=0, after=0)
p.paragraph_format.left_indent = Inches(0.18)
add_run(p, "ACS Division of Organic Chemistry Undergraduate Award  ·  GVSU Sustainability Champion Award", size=9, italic=True, color=GREY)

# ===================== RESEARCH =====================
section_header("Research & Translational Work")

entry_line("Northwestern University", " — PhD Candidate, Tyo Lab", "Sep 2023 – Present")
bullet_runs([("Developing data-integrated metabolic models of the gut microbiome to predict diet- and therapeutic-driven community dynamics — feeding directly into microbiome-targeted intervention design.", {})])

entry_line("Argonne National Laboratory", " — Assistant Computational Biologist, Henry Lab", "Jun 2021 – Present")
bullet_runs([("Core developer of ", {}), ("ModelSEED / ModelSEEDpy", {"bold": True}),
             (" — the most widely used open platform for genome-scale metabolic reconstruction (used by 100s of labs worldwide).", {})])
bullet_runs([("Curator of ", {}), ("openTECR", {"bold": True}),
             (", an open thermodynamics database underpinning rational metabolic engineering.", {})])

entry_line("Icahn School of Medicine at Mount Sinai", " — Bioinformatician, Karr Lab", "Dec 2020 – Apr 2022")
bullet_runs([("Built ", {}), ("BioSimulators", {"bold": True}), (" and ", {}), ("BioSimulations", {"bold": True}),
             (" — community-standard registries and a cloud platform enabling reproducible biological simulation across the field.", {})])

entry_line("Lawrence Berkeley National Lab", " — Software Engineer, Reactive Transport (Molins)", "Jun 2020 – Oct 2020")
bullet_runs([("Reactive-transport modeling of geochemical scaling in reverse-osmosis desalination — translated into peer-reviewed software and a pre-print.", {})])

entry_line("University of Victoria", " — Software Engineer, Buckley Lab (Green Safe Water)", "Jan 2020 – Apr 2022")
bullet_runs([("Simulations of bacterial biophysics & desalination geochemistry; first-author M.A.Sc. thesis.", {})])

entry_line("NC State & Shaw University", " — Metabolic Spectroscopist, Lucia Lab / GenoVerde BioSciences", "May 2019 – Dec 2019")
entry_line("Washington State University", " — NSF REU Fellow, Wolcott Lab (cellulose photocatalysis)", "Summer 2018")
entry_line("Grand Valley State University", " — Bioanalytical Chemist, Kovacs Lab (willow metabolomics)", "Oct 2016 – Jun 2019")

# ===================== PUBLICATIONS =====================
section_header("Selected Publications  (180 citations  ·  h-index 5  ·  full list at andrewfreiburger.com)")

def pub(authors_pre, name_bold, authors_post, title, venue, year):
    p = doc.add_paragraph()
    set_spacing(p, before=1, after=0)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    add_run(p, "• ")
    add_run(p, authors_pre)
    add_run(p, name_bold, bold=True)
    add_run(p, authors_post)
    add_run(p, f' "{title}." ')
    add_run(p, venue, italic=True)
    add_run(p, f", {year}.")

pub("Kennedy MS, ", "Freiburger A", ", Cooper M, et al.",
    "Diet outperforms microbial transplant to drive microbiome recovery in mice",
    "Nature", "2025")
pub("Borton MA, McGivern BB, …, ", "Freiburger A", ", et al.",
    "A functional microbiome catalogue crowdsourced from North American rivers",
    "Nature", "2025")
pub("Wood-Charlson EM, Henry CS, …, ", "Freiburger A", ", et al.",
    "KBase: Open-source platform for collaborative biological data analysis",
    "J. Molecular Biology", "2026")
pub("Shaikh B, Smith LP, …, ", "Freiburger AP", ", et al.",
    "BioSimulators: a central registry of simulation engines and services",
    "Nucleic Acids Research", "2022")
pub("Faria JP, Liu F, …, ", "Freiburger AP", ", et al.",
    "ModelSEED v2: High-throughput genome-scale metabolic model reconstruction",
    "bioRxiv", "2023")

# ===================== OPEN-SOURCE PLATFORMS =====================
section_header("Open-Source Platforms & Technical Skills")
bullet_runs([("Lead developer: ", {"bold": True}),
             ("CommScores", {"italic": True}), (" and ", {}), ("CommPhitting", {"italic": True}),
             (" — Python toolkits for scoring and fitting microbial community models.", {})])
bullet_runs([("Core contributor: ", {"bold": True}),
             ("ModelSEEDpy, BioSimulators, openTECR, KBase", {"italic": True}),
             (".  Stack: Python (COBRApy, pandas, scikit-learn), Jupyter, Git, Docker, HTML/CSS/JS, PHREEQC, MATLAB.", {})])

# ===================== ENTREPRENEURIAL TRAINING & LEADERSHIP =====================
section_header("Entrepreneurial Training & Leadership")
bullet_runs([("NSF Innovation Corps (I-Corps) Bay Area graduate", {"bold": True}),
             (" — customer-discovery and commercialization training at LBNL, 2020.", {})])
bullet_runs([("Co-organizer & logistics lead, ACS National Organic Symposium (NOS) 2025", {"bold": True}),
             (" (Northwestern host site); contributor to NOS 2027 planning.", {})])
bullet_runs([("Webmaster, ACS Division of Geochemistry", {"bold": True}),
             (" (2022–Present)  ·  ", {}),
             ("Asst. Webmaster, ACS Division of Organic Chemistry", {"bold": True}),
             (" (2020–Present)  ·  Instructor / TA across 7 university courses.", {})])

out = "/Users/andrewfreiburger/Documents/cv/Freiburger_Resume_1page.docx"
doc.save(out)
print(f"Saved: {out}")
